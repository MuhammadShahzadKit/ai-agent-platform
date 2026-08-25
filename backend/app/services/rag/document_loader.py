from __future__ import annotations

import csv
import io
from pathlib import Path

from langchain_core.documents import Document
from pypdf import PdfReader


# ============================================================
# SUPPORTED EXTENSIONS
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
    ".xls",
    ".pptx",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".bmp",
    ".tiff",
    ".tif",
}


# ============================================================
# COMMON METADATA
# ============================================================

def _base_metadata(
    path: Path,
    content_type: str,
) -> dict:

    return {
        "source": str(path),
        "filename": path.name,
        "file_type": path.suffix.lower(),
        "document_type": path.suffix.lower().lstrip("."),
        "content_type": content_type,
    }


# ============================================================
# TEXT CLEANING
# ============================================================

def _clean_text(
    text: str,
) -> str:

    if not text:
        return ""

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    lines = [
        line.rstrip()
        for line in text.split("\n")
    ]

    cleaned = []

    previous_blank = False

    for line in lines:

        if not line.strip():

            if previous_blank:
                continue

            previous_blank = True
            cleaned.append("")

        else:

            previous_blank = False
            cleaned.append(line)

    return "\n".join(
        cleaned
    ).strip()


# ============================================================
# PDF
# ============================================================

def _load_pdf(
    path: Path,
) -> list[Document]:

    reader = PdfReader(
        str(path)
    )

    documents = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        try:

            text = page.extract_text(
                extraction_mode="layout"
            )

        except TypeError:

            text = page.extract_text()

        text = _clean_text(
            text or ""
        )

        if not text:
            continue

        metadata = _base_metadata(
            path,
            "pdf_text",
        )

        metadata.update(
            {
                "page": page_number,
                "location_type": "page",
                "location": page_number,
            }
        )

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    return documents


# ============================================================
# DOCX
# ============================================================

def _load_docx(
    path: Path,
) -> list[Document]:

    from docx import Document as DocxDocument

    document = DocxDocument(
        str(path)
    )

    documents = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    paragraphs = []

    for paragraph in document.paragraphs:

        text = _clean_text(
            paragraph.text
        )

        if text:
            paragraphs.append(
                text
            )

    if paragraphs:

        text = "\n".join(
            paragraphs
        )

        metadata = _base_metadata(
            path,
            "docx_text",
        )

        metadata.update(
            {
                "location_type": "document",
                "location": 1,
                "section": "paragraphs",
            }
        )

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        rows = []

        for row in table.rows:

            cells = [
                _clean_text(
                    cell.text
                )
                for cell in row.cells
            ]

            rows.append(
                " | ".join(cells)
            )

        text = "\n".join(
            row
            for row in rows
            if row.strip()
        )

        if not text:
            continue

        metadata = _base_metadata(
            path,
            "docx_table",
        )

        metadata.update(
            {
                "location_type": "table",
                "location": table_index,
                "table": table_index,
            }
        )

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    return documents


# ============================================================
# PPTX
# ============================================================

def _load_pptx(
    path: Path,
) -> list[Document]:

    from pptx import Presentation

    presentation = Presentation(
        str(path)
    )

    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        parts = []

        for shape in slide.shapes:

            if not hasattr(
                shape,
                "text",
            ):
                continue

            text = _clean_text(
                shape.text
            )

            if text:
                parts.append(
                    text
                )

        text = "\n".join(
            parts
        ).strip()

        if not text:
            continue

        metadata = _base_metadata(
            path,
            "pptx_slide",
        )

        metadata.update(
            {
                "slide": slide_number,
                "location_type": "slide",
                "location": slide_number,
            }
        )

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    return documents


# ============================================================
# XLSX / XLS
# ============================================================

def _load_excel(
    path: Path,
) -> list[Document]:

    import pandas as pd

    documents = []

    workbook = pd.ExcelFile(
        path
    )

    for sheet_name in workbook.sheet_names:

        dataframe = pd.read_excel(
            path,
            sheet_name=sheet_name,
            dtype=str,
        )

        dataframe = dataframe.fillna(
            ""
        )

        rows = []

        for _, row in dataframe.iterrows():

            values = []

            for column in dataframe.columns:

                value = str(
                    row[column]
                ).strip()

                if value:
                    values.append(
                        f"{column}: {value}"
                    )

            if values:

                rows.append(
                    " | ".join(values)
                )

        text = "\n".join(
            rows
        ).strip()

        if not text:
            continue

        metadata = _base_metadata(
            path,
            "spreadsheet",
        )

        metadata.update(
            {
                "sheet": str(
                    sheet_name
                ),
                "location_type": "sheet",
                "location": str(
                    sheet_name
                ),
            }
        )

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    return documents


# ============================================================
# CSV
# ============================================================

def _load_csv(
    path: Path,
) -> list[Document]:

    raw = path.read_text(
        encoding="utf-8-sig",
        errors="replace",
    )

    sample = raw[:4096]

    try:

        dialect = csv.Sniffer().sniff(
            sample
        )

    except csv.Error:

        dialect = csv.excel

    reader = csv.reader(
        io.StringIO(raw),
        dialect,
    )

    rows = list(
        reader
    )

    if not rows:
        return []

    headers = rows[0]

    documents = []

    text_rows = []

    for row in rows[1:]:

        values = []

        for index, value in enumerate(
            row
        ):

            value = str(
                value
            ).strip()

            if not value:
                continue

            if index < len(headers):

                header = str(
                    headers[index]
                ).strip()

                if header:

                    values.append(
                        f"{header}: {value}"
                    )

                else:

                    values.append(
                        value
                    )

            else:

                values.append(
                    value
                )

        if values:

            text_rows.append(
                " | ".join(values)
            )

    if text_rows:

        text = "\n".join(
            text_rows
        )

        metadata = _base_metadata(
            path,
            "csv_table",
        )

        metadata.update(
            {
                "location_type": "table",
                "location": 1,
            }
        )

        documents.append(
            Document(
                page_content=text,
                metadata=metadata,
            )
        )

    return documents


# ============================================================
# TXT / MD
# ============================================================

def _load_text(
    path: Path,
) -> list[Document]:

    text = path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    text = _clean_text(
        text
    )

    if not text:
        return []

    metadata = _base_metadata(
        path,
        "plain_text",
    )

    metadata.update(
        {
            "location_type": "document",
            "location": 1,
        }
    )

    return [
        Document(
            page_content=text,
            metadata=metadata,
        )
    ]


# ============================================================
# IMAGE / OCR
# ============================================================

def _load_image(
    path: Path,
) -> list[Document]:

    try:

        from PIL import Image

    except ImportError:

        raise RuntimeError(
            "Pillow is required for image loading."
        )

    image = Image.open(
        path
    )

    # --------------------------------------------------------
    # Try OCR
    # --------------------------------------------------------

    try:

        import pytesseract

        # ----------------------------------------------------
        # Windows Tesseract configuration
        # ----------------------------------------------------

        tesseract_path = Path(
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        )

        if tesseract_path.exists():

            pytesseract.pytesseract.tesseract_cmd = str(
                tesseract_path
            )

        text = pytesseract.image_to_string(
            image
        )

    except ImportError:

        text = ""

    except Exception as exc:

        print(
            f"WARNING: OCR failed for "
            f"{path.name}: {exc}"
        )

        text = ""

    text = _clean_text(
        text
    )

    if not text:

        metadata = _base_metadata(
            path,
            "image",
        )

        metadata.update(
            {
                "location_type": "image",
                "location": 1,
                "ocr": False,
            }
        )

        return [
            Document(
                page_content=(
                    f"Image file: {path.name}"
                ),
                metadata=metadata,
            )
        ]

    metadata = _base_metadata(
        path,
        "image_ocr",
    )

    metadata.update(
        {
            "location_type": "image",
            "location": 1,
            "ocr": True,
        }
    )

    return [
        Document(
            page_content=text,
            metadata=metadata,
        )
    ]


# ============================================================
# MAIN LOADER
# ============================================================

def load_document(
    file_path: str,
) -> list[Document]:
    """
    Unified document loader.

    Supported:

        PDF
        DOCX
        PPTX
        XLSX
        XLS
        CSV
        TXT
        MD
        Images

    Returns LangChain Documents with normalized metadata.
    """

    path = Path(
        file_path
    )

    if not path.exists():

        raise FileNotFoundError(
            f"Document not found: {file_path}"
        )

    if not path.is_file():

        raise ValueError(
            f"Path is not a file: {file_path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    print(
        f"RAG LOADER: {path.name} "
        f"({extension})"
    )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        return _load_pdf(
            path
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == ".docx":

        return _load_docx(
            path
        )

    # --------------------------------------------------------
    # PPTX
    # --------------------------------------------------------

    if extension == ".pptx":

        return _load_pptx(
            path
        )

    # --------------------------------------------------------
    # Excel
    # --------------------------------------------------------

    if extension in {
        ".xlsx",
        ".xls",
    }:

        return _load_excel(
            path
        )

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    if extension == ".csv":

        return _load_csv(
            path
        )

    # --------------------------------------------------------
    # Text
    # --------------------------------------------------------

    if extension in {
        ".txt",
        ".md",
    }:

        return _load_text(
            path
        )

    # --------------------------------------------------------
    # Images
    # --------------------------------------------------------

    if extension in {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".bmp",
        ".tiff",
        ".tif",
    }:

        return _load_image(
            path
        )

    raise ValueError(
        f"No loader implemented for: {extension}"
    )